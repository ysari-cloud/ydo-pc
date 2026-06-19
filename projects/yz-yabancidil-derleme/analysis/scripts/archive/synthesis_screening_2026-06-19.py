#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Synthesis subset — ASSISTED FIRST-PASS eligibility screening + preliminary
deductive coding (C1-C8). Abstract-based, transparent, reproducible.
NOT a substitute for human full-text screening / Cohen's kappa: produces a
candidate pool + coding matrix for human verification (docs/04-kodlama-kitabi.md).
Reads focus raw files; writes outputs/synthesis_screening.xlsx.
"""
import os, re, csv, importlib.util
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
HERE=os.path.dirname(os.path.abspath(__file__)); ROOT=os.path.dirname(HERE)
RAW=os.path.join(ROOT,"raw_data"); OUT=os.path.join(ROOT,"outputs")
spec=importlib.util.spec_from_file_location("ba",os.path.join(HERE,"biblio_analysis.py"))
ba=importlib.util.module_from_spec(spec); spec.loader.exec_module(ba)

# ---- re-parse focus WITH abstracts (parser doesn't store AB) ----
def wos_focus():
    recs=[]; cur={}; tag=None
    for line in open(os.path.join(RAW,"wos_focus.txt"),encoding="utf-8-sig"):
        line=line.rstrip("\n")
        if line.startswith("ER"):
            if cur: recs.append(cur); cur={}
            tag=None; continue
        if line.startswith("   ") and tag: cur[tag].append(line.strip()); continue
        if len(line)>=2 and line[:2].strip(): tag=line[:2]; cur.setdefault(tag,[]).append(line[3:].strip())
    if cur: recs.append(cur)
    out=[]
    for r in recs:
        out.append({"db":"WoS","doi":ba.nd(r.get("DI",[""])[0] if r.get("DI") else ""),
            "title":" ".join(r.get("TI",[])),"year":(r.get("PY",[""])[0] if r.get("PY") else ""),
            "source":" ".join(r.get("SO",[])).strip(),"doctype":(r.get("DT",[""])[0] if r.get("DT") else ""),
            "abstract":" ".join(r.get("AB",[])),"keywords":";".join(r.get("DE",[])),
            "authors":[a.strip() for a in r.get("AF",[]) if a.strip()],
            "tc":int(re.sub(r"\D","",r.get("TC",["0"])[0]) or 0)})
    return out
def scopus_focus():
    out=[]
    for row in csv.DictReader(open(os.path.join(RAW,"scopus_focus.csv"),encoding="utf-8-sig")):
        out.append({"db":"Scopus","doi":ba.nd(row.get("DOI","")),"title":row.get("Title",""),
            "year":str(row.get("Year","")).strip(),"source":row.get("Source title","").strip(),
            "doctype":row.get("Document Type","").strip(),"abstract":row.get("Abstract",""),
            "keywords":row.get("Author Keywords",""),
            "authors":[a.strip() for a in row.get("Authors","").split(";") if a.strip()],
            "tc":int(row.get("Cited by","") or 0)})
    return out
def dedup(recs):
    sd=set(); st=set(); u=[]
    for r in recs:
        kd=r["doi"] or None; kt=(ba.nt(r["title"]),str(r["year"])[:4]) if r["title"] else None
        if (kd and kd in sd) or (kt and kt in st): continue
        if kd: sd.add(kd)
        if kt: st.add(kt)
        u.append(r)
    return u
focus=dedup(wos_focus()+scopus_focus())
print("focus unique:",len(focus),"| with abstract:",sum(1 for r in focus if len(r['abstract'])>40))

def blob(r): return (r["title"]+" "+r["abstract"]+" "+r["keywords"]).lower()
def has(text,terms): return any(t in text for t in terms)
def hits(text,terms): return sum(text.count(t) for t in terms)

# ---- lexicons ----
PROD=["generat","content creation","content generation","material","materials development",
      "instructional design","courseware","lesson plan","worksheet","automatic","automated",
      "item generation","test generation","question generation","authoring","creat","produc",
      "develop","design","tool","no-code","low-code"]
C1d=["app develop","application develop","website","web-based tool","courseware","software",
     "no-code","low-code","tool development","platform develop","build an app","chatbot develop"]
C1c=["co-design","co-create","assist teacher","support teacher","scaffold","suggestion","draft","co-author","teacher-ai"]
C1a=["content generation","content creation","generate material","generate content","produce material",
     "automatic generation","item generation","test generation","question generation","lesson plan","worksheet","generat"]
C1b=["chatbot","conversational agent","conversation","dialogue","speaking partner","practice partner",
     "interlocutor","oral practice","tutor","tutoring","interaction with"]
C2_teacher=["teacher-created","teacher created","teachers create","teachers produce","teacher-produced",
            "teachers design","teachers develop","teacher as designer","teacher-generated"]
C2_learner=["student-created","students create","learner-created","learners create","student-generated","learner-generated"]
C2_pub=["publisher","commercial material","expert-designed","textbook publisher"]
C3=["teacher role","teacher agency","professional identity","teacher identity","deskilling",
    "de-skilling","reskilling","upskilling","empowerment","teacher autonomy","teacher professional"]
C4=["alignment","curriculum","cefr","accessibilit","hallucinat","accuracy","quality assurance",
    "validation","pilot","expert review","reliab","valid","appropriateness"]
C5=["ai literacy","digital literacy","prompt","teacher education","teacher training",
    "professional development","tpack","competenc","ethic","skill set"]
C6=["labour","labor","platform depend","privacy","data protection","inequalit","digital divide",
    "equity","value capture","solutionism","surveillance","bias","critical"]
SKILLS={"reading":["reading"],"writing":["writing","essay","composition"],"listening":["listening"],
        "speaking":["speaking","oral","pronunciation"],"vocabulary":["vocabulary","lexical"],
        "grammar":["grammar","syntax"],"culture":["intercultural","culture"]}
EMP=["experiment","participant","survey","interview","quasi","randomi","mixed-method","mixed method",
     "qualitative","quantitative","data were","data was","sample of","n =","pre-test","post-test","questionnaire"]
CONC=["conceptual","framework","we argue","position paper","theoretical","perspective","viewpoint"]

def code_record(r):
    t=blob(r); c={}
    # eligibility
    is_journal = ("article" in r["doctype"].lower() or "review" in r["doctype"].lower()) and \
                 "conference" not in r["doctype"].lower() and "proceeding" not in r["doctype"].lower()
    is_excluded_type = any(x in r["doctype"].lower() for x in ["editorial","erratum","note","letter","retract"])
    prod_score = hits(t,PROD)
    has_abstract = len(r["abstract"])>40
    eligible = is_journal and not is_excluded_type and prod_score>=2 and has_abstract
    # C1 dominant role priority d>c>a>b
    if has(t,C1d): c["C1"]="C1d autonomous producer-developer"
    elif has(t,C1c): c["C1"]="C1c co-designer/assistant"
    elif has(t,C1a): c["C1"]="C1a content creator"
    elif has(t,C1b): c["C1"]="C1b conversational tutor"
    else: c["C1"]="?"
    # C2 locus
    loc=[]
    if has(t,C2_teacher): loc.append("teacher")
    if has(t,C2_learner): loc.append("learner")
    if has(t,C2_pub): loc.append("publisher")
    c["C2"]="+".join(loc) if loc else "?"
    c["C3"]="yes" if has(t,C3) else "-"
    c["C4"]="yes" if has(t,C4) else "-"
    c["C5"]="yes" if has(t,C5) else "-"
    c["C6"]="yes" if has(t,C6) else "-"
    sk=[s for s,kw in SKILLS.items() if has(t,kw)]
    c["C7_skills"]=",".join(sk) if sk else "general/unspec"
    # C8 claim/evidence type
    if "review" in r["doctype"].lower(): c["C8"]="review"
    elif has(t,EMP): c["C8"]="empirical"
    elif has(t,CONC): c["C8"]="conceptual"
    else: c["C8"]="unclear"
    c["prod_score"]=prod_score; c["eligible"]=eligible
    c["rel_rank"]=prod_score*1.0 + (0.02*r["tc"])   # relevance + minor citation weight
    return c

def slice_of(y):
    yi=int(y[:4]) if y[:4].isdigit() else 0
    return "2018-2021" if yi<=2021 and yi>=2018 else "2022-2023" if yi in (2022,2023) else "2024-2026" if yi>=2024 else "n/a"
for r in focus:
    r["code"]=code_record(r); r["slice"]=slice_of(r["year"])
elig=[r for r in focus if r["code"]["eligible"]]
elig.sort(key=lambda r:r["code"]["rel_rank"],reverse=True)

# ②: 40 = PRIORITY shortlist (NOT final). All 223 stay eligible candidates.
TARGET=40
priority=elig[:TARGET]
prio_ids=set(id(r) for r in priority)
# early-production bridge: eligible papers in 2018-2023 not already in priority
early=[r for r in elig if r["slice"] in ("2018-2021","2022-2023") and id(r) not in prio_ids]
early.sort(key=lambda r:(r["slice"],-r["code"]["rel_rank"]))
# combined human screening set, flag origin
for r in priority: r["set_origin"]="priority-40"
for r in early: r["set_origin"]="early-production bridge"
human_set=priority+early
from collections import Counter
elig_by_slice=Counter(r["slice"] for r in elig)
print("eligible (first-pass):",len(elig),"| by slice:",dict(elig_by_slice))
print("priority-40:",len(priority),"| early-production bridge added:",len(early),"| human screening set:",len(human_set))

# ---- Excel ----
HEAD=PatternFill("solid",fgColor="2B6CB0"); HF=Font(bold=True,color="FFFFFF")
YEL=PatternFill("solid",fgColor="FFF3CD"); GRN=PatternFill("solid",fgColor="E6F4EA"); BLU=PatternFill("solid",fgColor="DBEAFE")
thin=Side(style="thin",color="CBD5E0"); BORD=Border(thin,thin,thin,thin)
wb=Workbook()
def author_year(r):
    a=r["authors"][0].split(",")[0] if r["authors"] else "Anon"
    return f"{a} {r['year'][:4]}"
def add_sheet(name,records,title,show_origin=False):
    ws=wb.create_sheet(name)
    ws["A1"]=title; ws["A1"].font=Font(bold=True,size=12,color="1A365D")
    hdr=["#","Author-Year","Year","Slice","DB","Doc type","Source","Title","DOI","TC",
         "Prod.score","Eligible","C1 role","C2 locus","C3","C4","C5","C6","C7 skills","C8"]
    if show_origin: hdr.insert(2,"Set origin")
    hdr.append("Abstract (excerpt)")
    ws.append([]); ws.append(hdr)
    hr=ws.max_row
    for ci in range(1,len(hdr)+1):
        cc=ws.cell(hr,ci); cc.fill=HEAD; cc.font=HF; cc.alignment=Alignment(horizontal="center",wrap_text=True); cc.border=BORD
    for i,r in enumerate(records,1):
        c=r["code"]
        row=[i,author_year(r),r["year"][:4],r["slice"],r["db"],r["doctype"],r["source"][:40],
             r["title"][:90],r["doi"],r["tc"],c["prod_score"],"YES" if c["eligible"] else "no",
             c["C1"],c["C2"],c["C3"],c["C4"],c["C5"],c["C6"],c["C7_skills"],c["C8"]]
        if show_origin: row.insert(2,r.get("set_origin",""))
        row.append((r["abstract"][:300]+"…") if len(r["abstract"])>300 else r["abstract"])
        ws.append(row); rr=ws.max_row
        for ci in range(1,len(hdr)+1): ws.cell(rr,ci).border=BORD
        if show_origin and r.get("set_origin")=="early-production bridge":
            for ci in range(1,len(hdr)+1): ws.cell(rr,ci).fill=YEL
        if c["eligible"]: ws.cell(rr, (13 if show_origin else 12)).fill=GRN
    base=[4,18,7,10,8,22,30,50,26,7,10,9,26,18,6,6,6,6,18,12]
    if show_origin: base.insert(2,22)
    base.append(60)
    for i,w in enumerate(base,1): ws.column_dimensions[get_column_letter(i)].width=w
    ws.freeze_panes=ws.cell(hr+1,3)
    return ws

ov=wb.active; ov.title="ReadMe"
ov["A1"]="Sentez Alt-Kümesi — Özet Temelli Ön-Tarama ve C1–C8 Ön-Kodlama"
ov["A1"].font=Font(bold=True,size=13,color="1A365D")
notes=[ "",
 "AMAÇ: Odaklı 329 kaydı, tam-metin insan taramasına hazır bir aday havuza indirmek ve",
 "C1–C8 codebook'una göre ÖN (özet temelli) kodlama matrisi üretmek.",
 "",
 "ÇERÇEVELEME (kritik):",
 f" • Ön-uygun ADAY: {len(elig)} (dergi makale/derleme + üretim ilgisi prod_score≥2 + özet). HEPSİ aday kalır.",
 f" • ÖNCELİKLİ KISA LİSTE: ilk {len(priority)} (relevans+atıf sıralı). Bu NHAİ dahil-küme DEĞİLDİR;",
 "   tam-metin taraması için öncelik sırasıdır. Nihai n, tam-metin uygunluk + Cohen κ + 2. insan tarayıcı ile belirlenir.",
 f" • ERKEN-ÜRETİM KÖPRÜSÜ: 2018–2023 diliminde {len(early)} ön-uygun makale kısa listede değil; bunlar",
 "   'konuşma → üretim' kaymasının köprü kanıtıdır ve insan taramasına ÖNCELİKLİ-40 ile BİRLİKTE girer.",
 f" • İNSAN TARAMA SETİ (Human_screening_set sayfası): {len(priority)} + {len(early)} = {len(human_set)} makale.",
 "",
 "BÜTÜNLÜK: Bu otomatik ön-tarama, tam-metin uygunluk kararının ve κ'nın YERİNE GEÇMEZ.",
 "Etiketler özet/başlık/anahtar-sözcük temellidir; insan doğrulaması şarttır. κ/tam-metin UYDURULMADI.",
 "",
 f"Uygunluk by slice: 2018-2021={elig_by_slice['2018-2021']} · 2022-2023={elig_by_slice['2022-2023']} · 2024-2026={elig_by_slice['2024-2026']}",
 f"Fiili arama/export tarihi: 2026-06-19 (WoS DA alanı). Belge türü/yıl/dil sınırı export'ta uygulandı.",
 "",
 "C1 önceliği (codebook): d>c>a>b. '?' = özetten çıkmadı (insan bakmalı). Sarı satır = erken-üretim köprüsü.",
 "Sayfalar: Human_screening_set · Priority_40 · All_focus_scored (329) · Codebook.",
]
for n in notes: ov.append([n])
ov.column_dimensions["A"].width=112

add_sheet("Human_screening_set",human_set,f"İnsan Tam-Metin Tarama Seti = Öncelikli-40 + Erken-üretim köprüsü ({len(human_set)})",show_origin=True)
add_sheet("Priority_40",priority,f"Öncelikli Kısa Liste (ilk {len(priority)}; nihai DEĞİL — tam-metin önceliği)")
add_sheet("All_focus_scored",focus,f"Tüm Odaklı Kayıtlar — Skorlanmış ve Ön-Kodlanmış (n={len(focus)}; ön-uygun aday={len(elig)})")

cb=wb.create_sheet("Codebook")
cb["A1"]="C1–C8 Kodlama Kitabı (özet) — v0.1 (2026-06-19)"; cb["A1"].font=Font(bold=True,size=12,color="1A365D")
cbrows=[("C1 YZ rolü","C1a içerik üreticisi · C1b konuşma-öğretici · C1c eş-tasarımcı · C1d özerk üretici-geliştirici (öncelik d>c>a>b)"),
 ("C2 Üretim merkezi","yayıncı/uzman · öğretmen-üretici · öğrenen-üretici (fiili üretim öznesi)"),
 ("C3 Öğretmen rolü & failliği","tüketici↔geliştirici; kimlik; vasıfsızlaşma/yeniden vasıflanma"),
 ("C4 Pedagojik geçerlilik & KG","hizalama/CEFR · erişilebilirlik · hata/halüsinasyon · değerlendirme"),
 ("C5 Yeterlikler & öğretmen eğitimi","YZ/prompt okuryazarlığı · çıktı değerlendirme · etik · TPACK"),
 ("C6 Eleştirel/yapısal","emek · platform bağımlılığı · mahremiyet · eşitsizlik · değer yakalama · çözümcülük"),
 ("C7 Dil-özel","beceri (okuma/yazma/dinleme/konuşma/söz varlığı/dilbilgisi) · yön · düzey · dil"),
 ("C8 İddia/kanıt","kavramsal/empirik(nicel/nitel/karma)/derleme; kanıt gücü"),]
cb.append([]); cb.append(["Boyut","Kategoriler / karar kuralı"])
for a,bb in cbrows: cb.append([a,bb])
cb.column_dimensions["A"].width=28; cb.column_dimensions["B"].width=95

wb.save(os.path.join(OUT,"synthesis_screening.xlsx"))
print("Saved synthesis_screening.xlsx | sheets:",wb.sheetnames)
