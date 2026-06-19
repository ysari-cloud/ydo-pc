#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Assemble multi-sheet Excel from computed results (_results*.json)."""
import json, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
ROOT=os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT=os.path.join(ROOT,"outputs")
R=json.load(open(os.path.join(OUT,"_results.json"),encoding="utf-8"))
M=json.load(open(os.path.join(OUT,"_results_maps.json"),encoding="utf-8"))

HEAD=PatternFill("solid",fgColor="2B6CB0"); HF=Font(bold=True,color="FFFFFF")
TITLE=Font(bold=True,size=13,color="1A365D")
thin=Side(style="thin",color="CBD5E0"); BORD=Border(thin,thin,thin,thin)
wb=Workbook(); 

def sheet(name, title, headers, rows, widths=None):
    ws=wb.create_sheet(name)
    ws["A1"]=title; ws["A1"].font=TITLE; ws.append([])
    ws.append(headers)
    hr=ws.max_row
    for c in range(1,len(headers)+1):
        cell=ws.cell(hr,c); cell.fill=HEAD; cell.font=HF; cell.alignment=Alignment(horizontal="center")
        cell.border=BORD
    for row in rows:
        ws.append(row)
        for c in range(1,len(headers)+1): ws.cell(ws.max_row,c).border=BORD
    if widths:
        for i,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(i)].width=w
    ws.freeze_panes=ws.cell(hr+1,1)
    return ws

# Overview
ws=wb.active; ws.title="Overview"
ws["A1"]="GenAI × Foreign Language Education — Bibliometric Analysis"; ws["A1"].font=Font(bold=True,size=15,color="1A365D")
info=[("Run date","2026-06-19"),
      ("Databases","Web of Science Core Collection + Scopus"),
      ("Period","2018–2026 · English · Article/Review/Conference"),
      ("Broad corpus (unique)",R["n_broad"]),
      ("Focus corpus (unique)",R["n_focus"]),
      ("WoS docs w/ cited refs (maps)",M["wos_docs_with_CR"]),
      ("Tooling","Python (pandas/matplotlib/networkx); VOSviewer for maps"),
      ("Source integrity","All figures derive from raw exports; no fabricated data")]
ws.append([]); 
for k,v in info:
    ws.append([k,v]); ws.cell(ws.max_row,1).font=Font(bold=True)
ws.column_dimensions["A"].width=34; ws.column_dimensions["B"].width=70

sheet("AnnualProduction","Table 1. Annual scientific production",
      ["Year","Documents"], R["annual"], [12,14])
sheet("TopSources","Table 2. Most productive sources (Top 20)",
      ["Source","Documents"], R["sources"], [60,12])
sheet("TopAuthors","Table 3. Most productive authors (Top 20)",
      ["Author","Documents"], R["authors"], [34,12])
sheet("TopCountries","Table 4. Most productive countries (Top 20)",
      ["Country","Documents"], R["countries"], [26,12])
sheet("TopInstitutions","Table 5. Most productive institutions (Top 20)",
      ["Institution","Documents"], R["institutions"], [50,12])
sheet("MostCited","Table 6. Most cited documents (Top 20)",
      ["Title","Source","Year","Times cited","DB"], R["cited"], [70,34,8,12,8])
sheet("DocumentTypes","Table 7. Document type distribution",
      ["Document type","Count"], R["doctypes"], [34,12])
sheet("TopKeywords","Table 8. Most frequent author keywords (Top 40)",
      ["Author keyword","Documents"], R["keywords"], [40,12])
sheet("Bradford","Table 9. Bradford zones (source scattering)",
      ["Zone","# sources","# documents"], R["bradford_zones"], [10,14,14])
# Lotka
lot=[[npapers,nauth,round(frac*100,2)] for npapers,nauth,frac in R["lotka"]]
sheet("Lotka","Table 10. Author productivity (Lotka)",
      ["Papers per author","# authors","% of authors"], lot, [18,14,14])

# Thematic evolution
les=R["lexshare"]; rows=[]
for s,(n,conv,prod,cs,ps) in les.items():
    rows.append([s,n,conv,round(cs*100,1),prod,round(ps*100,1)])
sheet("ThematicShift","Table 11. Thematic evolution — conversation vs production framing",
      ["Time slice","Docs","Conv. docs","Conv. %","Prod. docs","Prod. %"], rows,[14,10,12,10,12,10])
# slice keywords
sk=R["slice_keywords"]; maxr=max(len(v) for v in sk.values())
rows=[]; slices=list(sk.keys())
for i in range(maxr):
    row=[]
    for s in slices:
        if i<len(sk[s]): row+= [sk[s][i][0], sk[s][i][1]]
        else: row+=["",""]
    rows.append(row)
hdr=[]; 
for s in slices: hdr+=[f"{s} keyword", "n"]
sheet("SliceKeywords","Table 12. Top keywords per time slice", hdr, rows,[28,6,28,6,28,6])

# network stats
ns=[["Keyword co-occurrence",R["net_keyword"]["nodes"],R["net_keyword"]["edges"],R["net_keyword"]["density"]],
    ["Co-citation (WoS)",M["cocitation"]["nodes"],M["cocitation"]["edges"],M["cocitation"]["density"]],
    ["Biblio. coupling countries (WoS)",M["biblio_coupling_countries"]["nodes"],M["biblio_coupling_countries"]["edges"],M["biblio_coupling_countries"]["density"]],
    ["Country collaboration",R["net_country"]["nodes"],R["net_country"]["edges"],R["net_country"]["density"]]]
sheet("NetworkStats","Table 13. Science-mapping network statistics",
      ["Network","Nodes","Edges","Density"], ns,[34,10,10,12])

wb.save(os.path.join(OUT,"bibliometric_tables.xlsx"))
print("Saved bibliometric_tables.xlsx with",len(wb.sheetnames),"sheets:",wb.sheetnames)
